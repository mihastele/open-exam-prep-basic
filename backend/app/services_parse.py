"""File → text extraction for uploads and photographed questions."""

import io
import json


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if name.endswith(".pptx"):
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [s.text.strip() for s in slide.shapes if getattr(s, "text", "").strip()]
            if texts:
                out.append(f"[Slide {i}]\n" + "\n".join(texts))
        return "\n\n".join(out).strip()
    if name.endswith((".txt", ".md", ".markdown")):
        return data.decode("utf-8", errors="replace").strip()
    if name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        return ocr_image(data)
    raise ValueError(f"Unsupported file type: {filename} (.pdf, .pptx, .txt, .md, image)")


def ocr_image(data: bytes) -> str:
    """OCR a photographed question. Needs system tesseract + pytesseract."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        raise RuntimeError(
            "Image OCR needs `pip install pytesseract pillow` plus the system "
            "`tesseract` binary — or type the question into the tutor instead."
        )
    img = Image.open(io.BytesIO(data))
    text = pytesseract.image_to_string(img).strip()
    if not text:
        raise RuntimeError("No text found in the image — try a sharper photo.")
    return text


def parse_json_response(raw: str):
    """Parse LLM JSON even when wrapped in ```json fences or prose."""
    text = raw.strip()
    if "```" in text:
        start = text.find("```")
        start = text.find("\n", start) + 1
        end = text.find("```", start)
        text = text[start:end if end != -1 else None].strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end != -1:
            return json.loads(text[start : end + 1])
        start, end = text.find("["), text.rfind("]")
        return json.loads(text[start : end + 1])
